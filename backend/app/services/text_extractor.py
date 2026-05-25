from io import BytesIO
from pathlib import Path

import pdfplumber
from docx import Document
from PyPDF2 import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


def extract_text_from_pdf(content: bytes) -> str:
    chunks: list[str] = []
    with pdfplumber.open(BytesIO(content)) as pdf:
        for page in pdf.pages:
            chunks.append(page.extract_text() or "")
    text = "\n".join(chunks).strip()
    if text:
        return text

    reader = PdfReader(BytesIO(content))
    return "\n".join(page.extract_text() or "" for page in reader.pages).strip()


def extract_text_from_docx(content: bytes) -> str:
    document = Document(BytesIO(content))
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" ".join(cell.text for cell in row.cells))
    return "\n".join(line for line in lines if line.strip())


def extract_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {suffix}")
    if suffix == ".pdf":
        return extract_text_from_pdf(content)
    if suffix == ".docx":
        return extract_text_from_docx(content)
    return content.decode("utf-8", errors="ignore")
